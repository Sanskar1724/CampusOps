"""API + PDF ingestion tests (isolated DB, TestClient)."""

from fastapi.testclient import TestClient

from backend.app import models
from backend.app.ingestion import pdf_source
from backend.app.main import app

client = TestClient(app, raise_server_exceptions=False)

def _build_pdf(payload_text: bytes) -> bytes:
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
         b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"),
        b"<< /Length %d >>\nstream\n" % len(payload_text) + payload_text + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref_at = len(out)
    out += b"xref\n0 %d\n" % (len(objs) + 1)
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += (b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
            % (len(objs) + 1, xref_at))
    return out


MINIMAL_PDF = _build_pdf(
    b"BT /F1 12 Tf 72 720 Td (DBMS mid-semester exam Monday 10am Room 405.) Tj ET")


def register(email="api@college.edu", prn="API1"):
    resp = client.post("/api/auth/register", json={
        "full_name": "Api User", "prn": prn, "department": "CE", "division": "A",
        "batch": "B1", "roll_number": "7", "semester": "5", "course": "B.Tech",
        "college_email": email, "password": "secret123"})
    assert resp.status_code == 200, resp.text
    return {"Authorization": f"Bearer {resp.json()['access_token']}"}


def test_register_login_and_profile(db_session):
    headers = register()
    assert client.get("/api/auth/me", headers=headers).status_code == 200
    bad = client.post("/api/auth/login",
                      json={"college_email": "api@college.edu", "password": "wrong"})
    assert bad.status_code == 401


def test_timetable_crud_and_chat(db_session):
    headers = register("tt@college.edu", "TT1")
    entry = {"day": 2, "subject": "DBMS", "start_time": "09:00",
             "end_time": "10:00", "room": "301"}
    assert client.post("/api/timetable/", json=entry, headers=headers).status_code == 200
    assert len(client.get("/api/timetable/", headers=headers).json()) == 1
    chat = client.post("/api/chat/", json={"text": "What is my next class?"},
                       headers=headers)
    assert chat.status_code == 200 and "DBMS" in chat.json()["reply"]


def test_deadline_task_flow(db_session):
    headers = register("plan@college.edu", "PL1")
    assert client.post("/api/planner/deadlines",
                       json={"title": "OS lab record", "due_at": "2026-09-12T17:00:00Z"},
                       headers=headers).status_code == 200
    assert client.post("/api/planner/tasks", json={"title": "Revise indexing"},
                       headers=headers).status_code == 200
    assert len(client.get("/api/planner/deadlines", headers=headers).json()) == 1


def test_pdf_upload_and_search(db_session):
    headers = register("doc@college.edu", "DOC1")
    client.post("/api/timetable/",
                json={"day": 0, "subject": "DBMS", "start_time": "09:00",
                      "end_time": "10:00", "room": "301"},
                headers=headers)
    resp = client.post("/api/documents/upload",
                       files={"file": ("calendar.pdf", MINIMAL_PDF, "application/pdf")},
                       headers=headers)
    assert resp.status_code == 200, resp.text
    assert resp.json()["facts"].get("subject") == "DBMS"
    hits = client.get("/api/documents/search", params={"q": "DBMS exam"},
                      headers=headers).json()
    assert hits and "DBMS" in hits[0]["text"]
    bad = client.post("/api/documents/upload",
                      files={"file": ("evil.exe", b"MZ", "application/octet-stream")},
                      headers=headers)
    assert bad.status_code == 400


def test_cross_student_access_blocked(db_session):
    alice = register("alice@college.edu", "AL1")
    bob = register("bob@college.edu", "BO1")
    entry_id = client.post(
        "/api/timetable/",
        json={"day": 0, "subject": "Secret", "start_time": "09:00", "end_time": "10:00"},
        headers=alice).json()["id"]
    assert client.delete(f"/api/timetable/{entry_id}", headers=bob).status_code == 404
    assert client.get("/api/timetable/", headers=bob).json() == []


def test_pdf_text_extraction_unit():
    text = pdf_source.extract_text("a.pdf", "application/pdf", MINIMAL_PDF)
    assert "DBMS" in text


def test_pdf_deep_scan_finds_all_facts():
    from datetime import datetime, timezone

    from backend.app.ingestion.extract import scan_facts
    facts = scan_facts(
        "DBMS assignment due 12/09. OS exam on 15/09 in Room 402. "
        "Holiday on 20/09. DBMS lecture moved from Room 301 to Room 405.",
        ["DBMS", "OS"], now=datetime(2026, 9, 9, tzinfo=timezone.utc))
    assert set(facts["subjects"]) == {"DBMS", "OS"}
    assert "2026-09-12" in facts["dates"] and "2026-09-15" in facts["dates"]
    assert facts["new_room"] == "405" and facts["old_room"] == "301"
    assert any(i["kind"] == "exam" for i in facts["items"])


def test_google_signin_creates_and_logs_in(db_session, monkeypatch):
    import backend.app.api.auth as auth_mod
    monkeypatch.setattr(auth_mod.email_source, "gmail_exchange_code",
                        lambda *a: {"access_token": "tok"})
    monkeypatch.setattr(auth_mod.email_source, "google_userinfo",
                        lambda t: {"email": "guser@college.edu", "name": "G User",
                                   "sub": "g123"})
    first = client.post("/api/auth/google/callback",
                        json={"code": "c", "redirect_uri": "http://x/cb"})
    assert first.status_code == 200, first.text
    me = client.get("/api/auth/me",
                    headers={"Authorization": f"Bearer {first.json()['access_token']}"})
    assert me.json()["onboarding_status"] == "pending"
    second = client.post("/api/auth/google/callback",
                         json={"code": "c", "redirect_uri": "http://x/cb"})
    assert second.status_code == 200  # existing user logs straight in


def test_gmail_get_callback_connects(db_session, monkeypatch):
    import backend.app.api.resources as res_mod
    from backend.app import security
    headers = register("gmailt@college.edu", "GM1")
    me = client.get("/api/auth/me", headers=headers).json()
    monkeypatch.setattr(res_mod.email_source, "gmail_exchange_code",
                        lambda *a: {"access_token": "tok", "refresh_token": "ref"})
    state = security.make_oauth_state(me["id"])
    resp = client.get("/api/integrations/gmail/callback",
                      params={"code": "c", "state": state})
    assert resp.status_code == 200 and "connected" in resp.text.lower()

    def _boom(self, db, student_id, max_results=20):
        raise ConnectionError("no network in tests")

    monkeypatch.setattr(res_mod.email_source.GmailSource, "fetch", _boom)
    sync = client.post("/api/email/sync", headers=headers)
    assert sync.status_code == 502  # fetch failure surfaces cleanly


def _doc_pdf_bytes(lines: list[str]) -> bytes:
    body = b"".join(b"BT /F1 12 Tf 72 %d Td (%s) Tj ET\n" % (720 - i * 20, line.encode())
                    for i, line in enumerate(lines))
    return _build_pdf(body)


def test_timetable_from_document_scan(db_session):
    headers = register("ttdoc@college.edu", "TTD1")
    pdf = _doc_pdf_bytes(["Monday 09:00-10:00 DBMS Room 301",
                          "Tuesday 11:00-12:00 OS Room 302"])
    resp = client.post("/api/timetable/from-document",
                       files={"file": ("tt.pdf", pdf, "application/pdf")},
                       headers=headers)
    assert resp.status_code == 200, resp.text
    assert resp.json()["entries"] == 2
    assert len(client.get("/api/timetable/", headers=headers).json()) == 2
    bad = client.post("/api/timetable/from-document",
                      files={"file": ("tt.pdf", MINIMAL_PDF, "application/pdf")},
                      headers=headers)
    assert bad.status_code == 400  # exam notice has no weekday+time rows


def test_system_status_board(db_session):
    headers = register("sys@college.edu", "SYS1")
    resp = client.get("/api/system/status", headers=headers)
    assert resp.status_code == 200
    names = {c["name"] for c in resp.json()["checks"]}
    assert {"database", "ai_model", "caspian_gateway", "telegram", "gmail"} <= names
    assert resp.json()["checks"][0]["ok"] is True  # database
    assert client.get("/api/system/status").status_code == 401  # auth required


def test_preferences_api_validation(db_session):
    headers = register("pref@college.edu", "PRF1")
    prefs = client.get("/api/student/preferences", headers=headers).json()
    assert prefs["notify_brief"] is True and prefs["hidden_subjects"] == []
    bad = client.put("/api/student/preferences",
                     json={"notify_channel": "pigeon"}, headers=headers)
    assert bad.status_code == 400
    bad2 = client.put("/api/student/preferences",
                      json={"quiet_start": 99}, headers=headers)
    assert bad2.status_code == 400
    ok = client.put("/api/student/preferences",
                    json={"notify_brief": False, "quiet_start": 22,
                          "quiet_end": 7, "notify_channel": "telegram"},
                    headers=headers).json()
    assert ok["updated"] == ["notify_brief", "notify_channel", "quiet_start", "quiet_end"]
    assert client.get("/api/student/preferences", headers=headers).json()["quiet_start"] == 22


def test_brief_respects_prefs_and_quiet_hours(db_session):
    from datetime import datetime, timezone
    from backend.app.jobs import run_daily_brief
    headers = register("quiet@college.edu", "QHT1")
    me = client.get("/api/auth/me", headers=headers).json()
    student = db_session.get(models.Student, me["id"])
    run_daily_brief(db_session, datetime(2026, 9, 9, 8, 0, tzinfo=timezone.utc))
    client.put("/api/student/preferences", json={"notify_brief": False}, headers=headers)
    before = db_session.query(models.Notification).filter_by(
        student_id=student.id, kind="brief").count()
    run_daily_brief(db_session, datetime(2026, 9, 10, 8, 0, tzinfo=timezone.utc))
    assert db_session.query(models.Notification).filter_by(
        student_id=student.id, kind="brief").count() == before
    client.put("/api/student/preferences",
               json={"notify_brief": True, "quiet_start": 0, "quiet_end": 23},
               headers=headers)
    run_daily_brief(db_session, datetime(2026, 9, 10, 8, 0, tzinfo=timezone.utc))
    assert db_session.query(models.Notification).filter_by(
        student_id=student.id, kind="brief").count() == before  # quiet all day


def test_delivery_suppressed_status(db_session):
    import backend.app.comms.proactive as pro_mod
    from backend.app.api.deps import set_pref
    headers = register("supp@college.edu", "SUP1")
    me = client.get("/api/auth/me", headers=headers).json()
    student = db_session.get(models.Student, me["id"])
    set_pref(db_session, student.id, "notify_brief", False)
    note = pro_mod.queue_notification(db_session, student.id, "brief", "Hi", "body")
    out = pro_mod.deliver_queued(db_session, note)
    assert out.status == "suppressed"


def test_telegram_commands_and_buttons(db_session):
    from backend.app.comms.service import COMMAND_TEXT, expand_command, quick_buttons
    assert expand_command("/today") == "What do I have today?"
    assert expand_command("/next@Sankiyy_bot") == "What is my next class?"
    assert expand_command("plain text") == "plain text"
    assert expand_command("/unknowncmd") == "/unknowncmd"
    buttons = quick_buttons()
    assert len(buttons) == 4  # thumb-sized: 4 main features only
    assert {b.data for b in buttons} == {"cmd:today", "cmd:next", "cmd:deadlines", "cmd:focus"}
    assert set(COMMAND_TEXT) >= {"today", "tomorrow", "next", "deadlines", "brief", "focus", "help"}


def test_on_action_rule_registered(db_session):
    from caspian import Action
    from caspian.core.commands import Host
    from backend.app.comms.client import build_caspian_app
    cx = build_caspian_app(mailbox="cmd-test", dispatch=False)
    assert len(cx.app.rules) == 2  # message + action rules
    interp = cx.interpret()
    event = Action(thread_id="t1", data="cmd:today", sender="s")  # type: ignore[arg-type]
    result = interp.run(cx.app, event, channel_name="email")
    assert any(isinstance(cmd, Host) for cmd in result.commands)


def test_help_endpoint(db_session):
    headers = register("help@college.edu", "HLP1")
    md = client.get("/api/system/help", headers=headers).json()["markdown"]
    assert "Timetable" in md and "Telegram" in md


def test_chat_threading_history(db_session):
    from backend.app import models
    headers = register("hist@college.edu", "HST1")
    me = client.get("/api/auth/me", headers=headers).json()
    client.post("/api/chat/", json={"text": "hi"}, headers=headers)
    client.post("/api/chat/", json={"text": "what is my next class?"}, headers=headers)
    conv = db_session.query(models.Conversation).filter_by(
        student_id=me["id"]).one()
    roles = [m.role for m in db_session.query(models.ChatMessage).filter_by(
        conversation_id=conv.id).order_by(models.ChatMessage.id)]
    assert roles == ["user", "agent", "user", "agent"]


def test_garbage_detector():
    from backend.app.ingestion.extract import garbage_score, is_garbage_text
    assert garbage_score("DBMS mid-semester exam Monday Room 405") < 0.3
    salad = "; n ; ; $ * Ess e rcc;E-tj t $ 9: j i F H *"
    assert garbage_score(salad * 10) > 0.6
    assert is_garbage_text(salad * 10, raw_bytes_len=50000) is True


def test_multiformat_extraction():
    from backend.app.ingestion.pdf_source import extract_any_text
    text, method = extract_any_text("n.txt", "text/plain", b"DBMS exam on 12/09")
    assert "DBMS" in text and method == "text"
    text, method = extract_any_text("n.csv", "text/csv", b"day,subject\nMon,DBMS")
    assert "DBMS" in text
    from openpyxl import Workbook
    import io as _io
    wb = Workbook()
    wb.active.append(["Subject", "Date"])
    wb.active.append(["OS", "15/09"])
    buf = _io.BytesIO()
    wb.save(buf)
    text, method = extract_any_text("n.xlsx",
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                    buf.getvalue())
    assert "OS" in text and method == "xlsx"
    from docx import Document as _Docx
    buf = _io.BytesIO()
    d = _Docx()
    d.add_paragraph("CN assignment due Friday")
    d.save(buf)
    text, method = extract_any_text(
        "n.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        buf.getvalue())
    assert "CN assignment" in text and method == "docx"


def test_scanned_pdf_uses_vision_ocr(db_session, monkeypatch):
    import backend.app.ingestion.pdf_source as pdf_mod
    monkeypatch.setattr(pdf_mod, "_layout_text", lambda data: "; n ; $ * Ess " * 60)
    monkeypatch.setattr(pdf_mod, "_ocr_pdf_pages",
                        lambda data: "Monday 09:00-10:00 DBMS Room 301")
    text, method = pdf_mod.extract_pdf_text(b"x" * 50000)
    assert method == "vision-ocr" and "DBMS" in text


def test_hybrid_search_skips_garbage(db_session):
    from backend.app import models
    from backend.app.memory import add_chunk, semantic_search
    headers = register("hyb@college.edu", "HYB1")
    me = client.get("/api/auth/me", headers=headers).json()
    doc = models.Document(student_id=me["id"], filename="real.pdf",
                          mime="application/pdf", size_bytes=10)
    db_session.add(doc)
    db_session.commit()
    db_session.refresh(doc)
    add_chunk(db_session, doc, 0, "DBMS mid-semester exam on Monday in Room 405")
    add_chunk(db_session, doc, 99, "; n ; $ * Ess " * 40)  # forced-in garbage
    hits = semantic_search(db_session, me["id"], "DBMS exam")
    assert hits and all("Ess" not in h["text"][:20] for h in hits)
    assert hits[0]["score"] > 0.3  # keyword recall lifts the real match


def test_gmail_refresh_on_401(db_session, monkeypatch):
    import backend.app.api.resources as res_mod
    headers = register("refr@college.edu", "REF1")
    me = client.get("/api/auth/me", headers=headers).json()
    from backend.app import models, security
    student = db_session.get(models.Student, me["id"])
    db_session.add(models.Integration(
        student_id=student.id, provider="gmail", status="connected",
        account_ref="refr@college.edu",
        meta_json='{"access_token": "stale", "refresh_token": "ref123"}'))
    db_session.commit()

    calls = {"fetch": 0}

    class FakeGmail:
        def __init__(self, token):
            self.token = token

        def fetch(self, db, student_id, max_results=20, days=7):
            calls["fetch"] += 1
            if self.token == "stale":
                raise Exception("401 Unauthorized for url")
            from backend.app.ingestion import RawItem
            return [RawItem(external_id="r:1", title="DBMS moved to Room 405",
                            sender="dept@college.edu",
                            body="DBMS lecture moved from Room 301 to Room 405.")]

    monkeypatch.setattr(res_mod.email_source, "GmailSource", FakeGmail)
    monkeypatch.setattr(res_mod.email_source, "gmail_refresh_access_token",
                        lambda *a: {"access_token": "fresh"})
    resp = client.post("/api/email/sync", headers=headers)
    assert resp.status_code == 200, resp.text
    assert resp.json()["processed"] == 1 and calls["fetch"] == 2
    db_session.refresh(student)
    integ = db_session.query(models.Integration).filter_by(
        student_id=student.id, provider="gmail").one()
    assert "fresh" in integ.meta_json


def test_proactive_telegram_fallback(db_session, monkeypatch):
    import backend.app.comms.proactive as pro_mod
    from backend.app import models
    headers = register("tgfb@college.edu", "TGFB1")
    me = client.get("/api/auth/me", headers=headers).json()
    student = db_session.get(models.Student, me["id"])
    student.telegram_chat_id = "777"
    db_session.commit()

    def boom_gateway(thread_id, text):
        raise Exception("SSL: CERTIFICATE_VERIFY_FAILED")

    sent = {}

    def fake_telegram(chat_id, text):
        sent["chat_id"] = chat_id
        assert "777" == chat_id

    monkeypatch.setattr(pro_mod, "_send_via_gateway", boom_gateway)
    monkeypatch.setattr(pro_mod, "_send_via_telegram", fake_telegram)
    note = pro_mod.notify_student(db_session, student.id, "brief", "Hi", "body")
    assert note.status == "sent" and sent["chat_id"] == "777"
    assert note.error == "" and note.sent_at is not None


def test_failed_delivery_keeps_body_intact(db_session, monkeypatch):
    import backend.app.comms.proactive as pro_mod
    headers = register("tgfail@college.edu", "TGF1")
    me = client.get("/api/auth/me", headers=headers).json()
    student = db_session.get(models.Student, me["id"])
    student.caspian_thread_id = "email:dead-thread"
    db_session.commit()
    monkeypatch.setattr(pro_mod, "_send_via_gateway",
                        lambda t, x: (_ for _ in ()).throw(Exception("down")))
    monkeypatch.setattr(pro_mod.config, "CASPIAN_API_KEY", "k")
    note = pro_mod.notify_student(db_session, student.id, "brief", "Hi", "body")
    assert note.status == "failed"
    assert note.body == "body" and "down" in note.error


def test_telegram_chat_id_captured(db_session):
    from backend.app.comms.handlers import _telegram_chat_id
    from caspian import Message
    assert _telegram_chat_id(Message(
        thread_id="t", text="hi", chat_kind="dm",  # type: ignore[arg-type]
        raw={"message": {"chat": {"id": 12345}}})) == "12345"
    assert _telegram_chat_id(Message(
        thread_id="t", text="hi", chat_kind="dm",  # type: ignore[arg-type]
        raw={"chat": {"id": 99}})) == "99"
    assert _telegram_chat_id(Message(
        thread_id="t", text="hi", chat_kind="dm")) == ""  # type: ignore[arg-type]


def test_timetable_scope_mine(db_session):
    headers = register("scope@college.edu", "SCP1")
    for row in ({"day": 0, "subject": "DBMS", "start_time": "09:00",
                 "end_time": "10:00", "room": "301", "division": "A", "batch": "B1"},
                {"day": 0, "subject": "Physics", "start_time": "09:00",
                 "end_time": "10:00", "room": "302", "division": "A", "batch": "B2"}):
        assert client.post("/api/timetable/", json=row, headers=headers).status_code == 200
    all_rows = client.get("/api/timetable/", headers=headers).json()
    mine = client.get("/api/timetable/", params={"scope": "mine"}, headers=headers).json()
    assert len(all_rows) == 2 and len(mine) == 1 and mine[0]["subject"] == "DBMS"


def test_rescan_rebuilds_chunks(db_session, monkeypatch):
    import backend.app.api.resources as res_mod
    headers = register("resc@college.edu", "RSC1")
    me = client.get("/api/auth/me", headers=headers).json()
    from backend.app import models
    doc = models.Document(student_id=me["id"], filename="old.pdf",
                          mime="application/pdf", size_bytes=10, facts_json="{}")
    db_session.add(doc)
    db_session.commit()
    db_session.refresh(doc)
    from backend.app.memory import add_chunk
    add_chunk(db_session, doc, 0, "; n ; $ * Ess " * 40)
    monkeypatch.setattr(res_mod.pdf_source, "extract_any_text",
                        lambda fn, mime, data: ("OS exam on 15/09 in Room 402", "vision-ocr"))
    resp = client.post(f"/api/documents/{doc.id}/rescan",
                       files={"file": ("old.pdf", b"%PDF-fake", "application/pdf")},
                       headers=headers)
    assert resp.status_code == 200, resp.text
    assert resp.json()["extraction"] == "vision-ocr"
    texts = [c.text for c in
             db_session.query(models.DocumentChunk).filter_by(document_id=doc.id)]
    assert texts and all("Ess" not in t for t in texts)
